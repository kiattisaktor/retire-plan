import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
wb.remove(wb.active)

NAVY_HEADER = "1B365D"
EMERALD_HEADER = "0D5C3A"
BLUE_HEADER = "1E3A8A"
GOLD_ACCENT = "D4AF37"
DARK_TEXT = "1A1A1A"
LIGHT_BG = "F8FAFC"
ZEBRA_BG = "F1F5F9"
BORDER_GRAY = "CBD5E1"
WHITE = "FFFFFF"
ACCENT_BLUE = "E0F2FE"
ACCENT_GREEN = "DCFCE7"

header_font = Font(name="Arial", size=10, bold=True, color=WHITE)
title_font = Font(name="Arial", size=16, bold=True, color=NAVY_HEADER)
subtitle_font = Font(name="Arial", size=11, italic=True, color="64748B")
bold_font = Font(name="Arial", size=10, bold=True, color=DARK_TEXT)
normal_font = Font(name="Arial", size=10, color=DARK_TEXT)

thin_border = Border(
    left=Side(style="thin", color=BORDER_GRAY),
    right=Side(style="thin", color=BORDER_GRAY),
    top=Side(style="thin", color=BORDER_GRAY),
    bottom=Side(style="thin", color=BORDER_GRAY)
)
total_border = Border(
    left=Side(style="thin", color=BORDER_GRAY),
    right=Side(style="thin", color=BORDER_GRAY),
    top=Side(style="thin", color=NAVY_HEADER),
    bottom=Side(style="double", color=NAVY_HEADER)
)

# -------------------------------------------------------------
# SHEET 1: แผนสะสมเงิน 38-60 ปี (Accumulation Phase)
# -------------------------------------------------------------
ws1 = wb.create_sheet(title="1_แผนสะสมเงิน 38-60 ปี")
ws1.views.sheetView[0].showGridLines = True

ws1["B2"] = "แผนการออม การลงทุน และรายได้ สู่เป้าหมายเกษียณ 25 ล้านบาท"
ws1["B2"].font = title_font
ws1["B3"] = "พนักงาน กฟผ. ช่วงอายุ 38 - 60 ปี (รวม 23 ปี) | คำนวณรายได้จากเงินเดือน + โบนัส 2 เดือน"
ws1["B3"].font = subtitle_font

headers1 = [
    ("ลำดับ", 6),
    ("อายุ (ปี)", 10),
    ("ฐานเงินเดือน (บาท)", 16),
    ("โบนัสประจำปี (~2 ด.)", 18),
    ("รวมรายได้ต่อปี (บาท)", 20),
    ("PVD สะสม 15% (บ./ด.)", 18),
    ("หุ้น สอ.กฟผ. (บ./ด.)", 18),
    ("กองทุน RMF (บ./ด.)", 16),
    ("ประกันเดิม 3 ฉบับ (บ./ด.)", 18),
    ("ฝากพิเศษ สอ. (บ./ด.)", 16),
    ("รวมหักเงินเดือนประจำ", 20),
    ("บำนาญใหม่ (โบนัส)", 18),
    ("รวมจ่ายทั้งปี (บาท)", 20),
    ("เงินเหลือใช้สุทธิ/ปี", 20),
    ("หมุดหมายสำคัญ / Strategic Action", 40)
]

start_row = 5
for col_idx, (h_text, width) in enumerate(headers1, start=2):
    cell = ws1.cell(row=start_row, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.column_dimensions[get_column_letter(col_idx)].width = width

ws1.row_dimensions[start_row].height = 36

raw_data_s1 = [
    (1, 38, 63000, 20000, 2000, 9583, 0, 0, "เริ่มแผน: สอ. 20k/ด. ล็อกเป้า 4M, RMF 2k/ด., PVD 15%"),
    (2, 39, 66150, 20000, 2000, 9583, 0, 0, "ฐานเงินเดือนโต 5% ต่อปี"),
    (3, 40, 69458, 20000, 2000, 9583, 0, 0, "สะสมหุ้น สอ. ต่อเนื่อง"),
    (4, 41, 72930, 20000, 2000, 10000, 0, 0, "ปรับเบี้ยประกันสุขภาพตามช่วงอายุ"),
    (5, 42, 76577, 20000, 2000, 10000, 0, 0, "ปีสุดท้ายที่ส่ง สอ.กฟผ. ครบเป้าหมายเงินต้น 4,000,000 บ."),
    (6, 43, 80406, 0, 24000, 10000, 0, 0, "หยุดส่ง สอ. (ปันผล 15k/ด. โอนให้แม่ 100%), เร่ง RMF เป็น 24k/ด."),
    (7, 44, 84426, 0, 24000, 10000, 0, 0, "เร่งสปีดหุ้นโลกผ่าน PVD 15% + RMF 24k"),
    (8, 45, 88647, 0, 24000, 10000, 0, 0, "ช่วงสะสมสินทรัพย์เติบโตสูง"),
    (9, 46, 93080, 0, 24000, 10667, 0, 0, "ปรับเบี้ยประกันสุขภาพตามช่วงอายุ 46-50 ปี"),
    (10, 47, 97734, 0, 24000, 10667, 0, 0, "เตรียมแตะเพดานเงินเดือน กฟผ."),
    (11, 48, 100000, 0, 24000, 10667, 0, 0, "ฐานเงินเดือนแตะเพดาน 100,000 บ. (PVD หักคงที่ 15,000 บ./ด.)"),
    (12, 49, 100000, 0, 24000, 10667, 0, 0, "ปีสุดท้ายของการเร่ง RMF 24k/ด."),
    (13, 50, 100000, 0, 16000, 10667, 0, 100000, "RMF ลดเหลือ 16k/ด. + ซื้อบำนาญใหม่ #3 จ่าย 100k/ปี จากโบนัส (ปี 1/5)"),
    (14, 51, 100000, 0, 16000, 11500, 0, 100000, "บำนาญใหม่ #3 จ่าย 100k/ปี (ปี 2/5) | ปรับเบี้ยสุขภาพช่วง 51-55"),
    (15, 52, 100000, 0, 16000, 11500, 0, 100000, "บำนาญใหม่ #3 จ่าย 100k/ปี (ปี 3/5)"),
    (16, 53, 100000, 0, 16000, 11500, 0, 200000, "ซื้อบำนาญ 85/5 เพิ่ม 100k/ปี (ปี 1/5) ซ้อนบำนาญ #3 (ปี 4/5) รวม 200k/ปี"),
    (17, 54, 100000, 0, 16000, 11500, 0, 200000, "บำนาญ 85/5 (ปี 2/5) + บำนาญ #3 (ปี 5/5 งวดสุดท้าย) รวม 200k/ปี"),
    (18, 55, 100000, 0, 16000, 11500, 0, 100000, "สับ PVD+RMF เข้าตราสารหนี้ | บำนาญ 85/5 (ปี 3/5)"),
    (19, 56, 100000, 0, 16000, 12500, 0, 100000, "บำนาญ 85/5 (ปี 4/5) | ปรับเบี้ยสุขภาพช่วง 56-60"),
    (20, 57, 100000, 0, 16000, 12500, 0, 100000, "ส่งประกันสะสมทรัพย์งวดสุดท้าย | บำนาญ 85/5 งวดสุดท้าย (ปี 5/5)"),
    (21, 58, 100000, 0, 16000, 9167, 0, 0, "หยุดส่งสะสมทรัพย์ (รับเงินคืน 250k เข้า e-Savings) เหลือส่งบำนาญเดิม+สุขภาพ"),
    (22, 59, 100000, 0, 16000, 9167, 0, 0, "เตรียมความพร้อมก่อนเกษียณ 1 ปี (วางแผนโยก PVD เข้า สอ. ภายใน 60 วัน)"),
    (23, 60, 100000, 0, 16000, 9167, 0, 0, "ปีสุดท้าย: ส่งบำนาญเดิมงวดสุดท้าย + รับเงินชดเชยเกษียณ 400 วัน 1.312M + โยก PVD 5M ฝากพิเศษเกษียณสุข")
]

for row_idx, r in enumerate(raw_data_s1, start=6):
    ws1.row_dimensions[row_idx].height = 22
    bg_color = ZEBRA_BG if row_idx % 2 == 0 else WHITE
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    
    seq, age, salary, coop, rmf, ins_orig, coop_spec, bonus_pen, note = r
    
    # Check milestone highlight
    is_milestone = age in [42, 43, 48, 53, 57, 60]
    if is_milestone:
        fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    
    # Col B: Seq
    c = ws1.cell(row=row_idx, column=2, value=seq)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col C: Age
    c = ws1.cell(row=row_idx, column=3, value=age)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.font = Font(name="Arial", size=10, bold=True, color=EMERALD_HEADER if is_milestone else DARK_TEXT)
    c.fill = fill; c.border = thin_border
    
    # Col D: Salary
    c = ws1.cell(row=row_idx, column=4, value=salary)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col E: Bonus (~2 months)
    c = ws1.cell(row=row_idx, column=5, value=f"=D{row_idx}*2")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = Font(name="Arial", size=10, color="2563EB")
    c.fill = fill; c.border = thin_border
    
    # Col F: Total Annual Income = Salary*12 + Bonus
    c = ws1.cell(row=row_idx, column=6, value=f"=D{row_idx}*12+E{row_idx}")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = Font(name="Arial", size=10, bold=True, color="1E40AF")
    c.fill = fill; c.border = thin_border
    
    # Col G: PVD 15%
    c = ws1.cell(row=row_idx, column=7, value=f"=ROUND(D{row_idx}*0.15, 0)")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col H: Coop
    c = ws1.cell(row=row_idx, column=8, value=coop)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col I: RMF
    c = ws1.cell(row=row_idx, column=9, value=rmf)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col J: Ins Original
    c = ws1.cell(row=row_idx, column=10, value=ins_orig)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col K: Coop Special
    c = ws1.cell(row=row_idx, column=11, value=coop_spec)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col L: Total Monthly Deductions = SUM(G:K)
    c = ws1.cell(row=row_idx, column=12, value=f"=SUM(G{row_idx}:K{row_idx})")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = bold_font; c.fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
    c.border = thin_border
    
    # Col M: Bonus Pension Annual
    c = ws1.cell(row=row_idx, column=13, value=bonus_pen)
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    # Col N: Total Annual Payment = L*12 + M
    c = ws1.cell(row=row_idx, column=14, value=f"=L{row_idx}*12+M{row_idx}")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = Font(name="Arial", size=10, bold=True, color=EMERALD_HEADER)
    c.fill = PatternFill(start_color=ACCENT_GREEN, end_color=ACCENT_GREEN, fill_type="solid")
    c.border = thin_border
    
    # Col O: Net Disposable Income = F - N
    c = ws1.cell(row=row_idx, column=15, value=f"=F{row_idx}-N{row_idx}")
    c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = Font(name="Arial", size=10, bold=True, color="047857")
    c.fill = fill; c.border = thin_border
    
    # Col P: Notes
    c = ws1.cell(row=row_idx, column=16, value=note)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border

# Total Row
tot_row = 6 + len(raw_data_s1)
ws1.row_dimensions[tot_row].height = 26
ws1.cell(row=tot_row, column=2, value="").border = total_border
ws1.cell(row=tot_row, column=3, value="รวมตลอด 23 ปี").font = bold_font
ws1.cell(row=tot_row, column=3).alignment = Alignment(horizontal="center", vertical="center")
ws1.cell(row=tot_row, column=3).border = total_border

ws1.cell(row=tot_row, column=4, value=f"=SUM(D6:D{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=5, value=f"=SUM(E6:E{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=6, value=f"=SUM(F6:F{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=7, value=f"=SUM(G6:G{tot_row-1})*12").number_format = "#,##0"
ws1.cell(row=tot_row, column=8, value=f"=SUM(H6:H{tot_row-1})*12").number_format = "#,##0"
ws1.cell(row=tot_row, column=9, value=f"=SUM(I6:I{tot_row-1})*12").number_format = "#,##0"
ws1.cell(row=tot_row, column=10, value=f"=SUM(J6:J{tot_row-1})*12").number_format = "#,##0"
ws1.cell(row=tot_row, column=11, value=f"=SUM(K6:K{tot_row-1})*12").number_format = "#,##0"
ws1.cell(row=tot_row, column=12, value="-").alignment = Alignment(horizontal="center", vertical="center")
ws1.cell(row=tot_row, column=13, value=f"=SUM(M6:M{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=14, value=f"=SUM(N6:N{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=15, value=f"=SUM(O6:O{tot_row-1})").number_format = "#,##0"
ws1.cell(row=tot_row, column=16, value="ยอดรวมสะสมตลอด 23 ปี สู่เป้าหมาย 25 ล้านบาท").font = Font(name="Arial", size=10, italic=True)

for col in range(2, 17):
    cell = ws1.cell(row=tot_row, column=col)
    cell.font = bold_font
    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    cell.border = total_border
    if col in [4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15]:
        cell.alignment = Alignment(horizontal="right", vertical="center")

# Milestone Notes section below the table in Excel
note_start = tot_row + 2
ws1.cell(row=note_start, column=3, value="📌 บันทึกความหมายของปีหมุดหมายสำคัญ (แถบไฮไลท์สีเขียว):").font = Font(name="Arial", size=11, bold=True, color=EMERALD_HEADER)
milestone_notes_excel = [
    ("อายุ 42 ปี", "ส่งหุ้น สอ.กฟผ. ปีสุดท้าย สะสมเงินต้นครบ 4,000,000 บาท ตามเป้าหมาย"),
    ("อายุ 43 ปี", "จุดเปลี่ยนกลยุทธ์ (Pivot Year) หยุดส่งหุ้น สอ. นำปันผล 15k/ด. โอนให้แม่ 100% และเร่ง RMF เป็น 24,000 บ./ด."),
    ("อายุ 48 ปี", "ฐานเงินเดือน กฟผ. แตะเพดาน 100,000 บาท ทำให้ PVD 15% ล็อกสูงสุด 15,000 บ./ด. คงที่ตลอดไป"),
    ("อายุ 53 ปี", "ภาระบำนาญชนยอดสูงสุด ซื้อบำนาญ 85/5 ซ้อนกับบำนาญ #3 หักจากเงินโบนัสรวมสูงสุด 200,000 บ./ปี"),
    ("อายุ 57 ปี", "ส่งเบี้ยสะสมทรัพย์และบำนาญ 85/5 งวดสุดท้าย ปลดภาระจ่ายประกัน (เตรียมรับเงินคืน 250k ตอนอายุ 58 และ 62)"),
    ("อายุ 60 ปี", "วันเกษียณอายุ กฟผ. ส่งบำนาญเดิมงวดสุดท้าย + รับเงินชดเชยเกษียณ 400 วัน (~1.312M สุทธิ) สินทรัพย์แตะ 25 ล้านบาท")
]

for idx, (yr, desc) in enumerate(milestone_notes_excel, start=note_start+1):
    ws1.cell(row=idx, column=3, value=yr).font = bold_font
    ws1.cell(row=idx, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=idx, column=4, value=desc).font = normal_font

# -------------------------------------------------------------
# SHEET 2: โครงสร้าง 6 ตะกร้า 25 ล้าน
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="2_โครงสร้าง 6 ตะกร้า 25 ล้าน")
ws2.views.sheetView[0].showGridLines = True

ws2["B2"] = "โครงสร้างพอร์ตสินทรัพย์ 25.00 ล้านบาท ณ วันเกษียณอายุ 60 ปี"
ws2["B2"].font = title_font
ws2["B3"] = "กระจายความเสี่ยง 6 ตะกร้าสินทรัพย์ + ทองคำแท่ง 5 บาท สำรองพิเศษ | ผลิตกระแสเงินสดมั่นคง"
ws2["B3"].font = subtitle_font

headers2 = [
    ("ตะกร้าที่", 10),
    ("รายการสินทรัพย์ / สถาบัน", 32),
    ("เงินต้นวันเกษียณ (บาท)", 22),
    ("สัดส่วน (%)", 14),
    ("ผลตอบแทนสุทธิ (%/ปี)", 20),
    ("กระแสเงินสดรับ (บาท/ปี)", 22),
    ("กระแสเงินสดรับ (บาท/เดือน)", 24),
    ("บทบาทและหน้าที่ในแผนเกษียณ", 42)
]

start_row_2 = 5
for col_idx, (h_text, width) in enumerate(headers2, start=2):
    cell = ws2.cell(row=start_row_2, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = PatternFill(start_color=EMERALD_HEADER, end_color=EMERALD_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.column_dimensions[get_column_letter(col_idx)].width = width

ws2.row_dimensions[start_row_2].height = 36

baskets_data = [
    (1, "เงินฝากพิเศษเกษียณสุข สอ.กฟผ.", 5000000, 0.0350, "ปลอดภาษี", "โยกจาก PVD สูงสุด 5 ล้านบาท ภายใน 60 วัน"),
    (2, "ทุนเรือนหุ้น สอ.กฟผ. (ท่อคุณแม่)", 4000000, 0.0450, "ปลอดภาษี", "โอนตรงให้คุณแม่ 100% แยกบัญชีขาด (15k/ด.)"),
    (3, "พันธบัตรรัฐบาล / วอลเล็ต สบม.", 6500000, 0.02295, "หลังหักภาษี 15%", "ล็อกผลตอบแทนระยะยาว ไร้ความเสี่ยงผิดนัด (เพิ่ม 1 ล้าน)"),
    (4, "กองทุนตลาดเงิน Treasury MMF", 4000000, 0.0190, "ปลอดภาษี", "เสริมสภาพคล่องสูง T+1 พร้อมสับเปลี่ยน"),
    (5, "สลากออมทรัพย์ 3 สถาบัน", 3000000, 0.0120, "ดอกเบี้ย+รางวัลเฉลี่ย", "ทุนสำรองสภาพคล่องด่านแรก ทยอยดึงใช้หลังอายุ 70"),
    (6, "e-Savings + คืนประกัน + เงิน 400 วัน", 2500000, 0.01275, "หลังหักภาษี 15%", "สภาพคล่องฉุกเฉิน ดูแลสุขภาพและกองทุนคุณแม่")
]

for idx, b in enumerate(baskets_data, start=6):
    ws2.row_dimensions[idx].height = 24
    bg_color = ZEBRA_BG if idx % 2 == 0 else WHITE
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    b_no, name, principal, rate, tax_info, role = b
    
    ws2.cell(row=idx, column=2, value=b_no).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=idx, column=2).font = bold_font; ws2.cell(row=idx, column=2).fill = fill; ws2.cell(row=idx, column=2).border = thin_border
    
    ws2.cell(row=idx, column=3, value=name).font = bold_font; ws2.cell(row=idx, column=3).fill = fill; ws2.cell(row=idx, column=3).border = thin_border
    
    c = ws2.cell(row=idx, column=4, value=principal); c.number_format = "#,##0"; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = bold_font; c.fill = fill; c.border = thin_border
    
    c = ws2.cell(row=idx, column=5, value=f"=D{idx}/D12"); c.number_format = "0.0%"; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    c = ws2.cell(row=idx, column=6, value=rate); c.number_format = "0.000%"; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    c = ws2.cell(row=idx, column=7, value=f"=ROUND(D{idx}*F{idx}, 0)"); c.number_format = "#,##0"; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = normal_font; c.fill = fill; c.border = thin_border
    
    c = ws2.cell(row=idx, column=8, value=f"=ROUND(G{idx}/12, 0)"); c.number_format = "#,##0"; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = Font(name="Arial", size=10, bold=True, color=EMERALD_HEADER); c.fill = PatternFill(start_color=ACCENT_GREEN, end_color=ACCENT_GREEN, fill_type="solid"); c.border = thin_border
    
    c = ws2.cell(row=idx, column=9, value=f"{role} ({tax_info})"); c.font = normal_font; c.fill = fill; c.border = thin_border

# Total Row
tot_b_row = 12
ws2.row_dimensions[tot_b_row].height = 26
ws2.cell(row=tot_b_row, column=2, value="").border = total_border
ws2.cell(row=tot_b_row, column=3, value="รวมพอร์ตสินทรัพย์ 6 ตะกร้า").font = bold_font
ws2.cell(row=tot_b_row, column=3).alignment = Alignment(horizontal="center", vertical="center"); ws2.cell(row=tot_b_row, column=3).border = total_border

ws2.cell(row=tot_b_row, column=4, value=f"=SUM(D6:D11)").number_format = "#,##0"
ws2.cell(row=tot_b_row, column=5, value=f"=SUM(E6:E11)").number_format = "0.0%"
ws2.cell(row=tot_b_row, column=6, value=f"=G12/D12").number_format = "0.000%"
ws2.cell(row=tot_b_row, column=7, value=f"=SUM(G6:G11)").number_format = "#,##0"
ws2.cell(row=tot_b_row, column=8, value=f"=SUM(H6:H11)").number_format = "#,##0"
ws2.cell(row=tot_b_row, column=9, value="ผลิตกระแสเงินสดเฉลี่ย 54,004 บ./เดือน (รวมท่อคุณแม่ 15,000 บ./ด.)").font = Font(name="Arial", size=10, italic=True)

for col in range(2, 10):
    cell = ws2.cell(row=tot_b_row, column=col)
    cell.font = bold_font; cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"); cell.border = total_border
    if col in [4, 5, 6, 7, 8]:
        cell.alignment = Alignment(horizontal="right", vertical="center")

# Gold buffer row
gold_row = 14
ws2.cell(row=gold_row, column=2, value="★").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=gold_row, column=3, value="ทองคำแท่งบริสุทธิ์ (96.5% / 99.99%)").font = bold_font
ws2.cell(row=gold_row, column=4, value="5 บาททองคำ").font = bold_font
ws2.cell(row=gold_row, column=4).alignment = Alignment(horizontal="right", vertical="center")
ws2.cell(row=gold_row, column=5, value="-").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=gold_row, column=6, value="-").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=gold_row, column=7, value="-").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=gold_row, column=8, value="-").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=gold_row, column=9, value="Wealth Buffer สินทรัพย์สำรองพิเศษฉุกเฉิน ถือครองแยก 100%").font = Font(name="Arial", size=10, italic=True)

for col in range(2, 10):
    cell = ws2.cell(row=gold_row, column=col)
    cell.fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid"); cell.border = thin_border

# -------------------------------------------------------------
# SHEET 3: กระแสเงินสดหลังเกษียณ (Dual Cash Flow & Drawdown)
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="3_กระแสเงินสดหลังเกษียณ")
ws3.views.sheetView[0].showGridLines = True

ws3["B2"] = "โครงสร้างกระแสเงินสดรับ-จ่ายหลังเกษียณอายุ 60 ปี (Dual Pipelines)"
ws3["B2"].font = title_font
ws3["B3"] = "ท่อที่ 1 คุณแม่ (15,000 บ./ด.) | ท่อที่ 2 ส่วนตัวคุณ (50,038 บ./ด.) พร้อมแผน Drawdown ปลอดภัย"
ws3["B3"].font = subtitle_font

ws3["B5"] = "1. ท่อกระแสเงินสดรายเดือนหลังอายุ 60 ปี (Dual Cash Flow Pipelines)"
ws3["B5"].font = Font(name="Arial", size=12, bold=True, color=NAVY_HEADER)

headers3_1 = [
    ("ท่อกระแสเงินสด", 26),
    ("แหล่งที่มาของเงิน", 34),
    ("รายรับต่อปี (บาท)", 20),
    ("รายรับต่อเดือน (บาท)", 22),
    ("การจัดสรรและวัตถุประสงค์", 40)
]
for col_idx, (h_text, width) in enumerate(headers3_1, start=2):
    cell = ws3.cell(row=6, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws3.column_dimensions[get_column_letter(col_idx)].width = width
ws3.row_dimensions[6].height = 28

cf_pipeline_data = [
    ("ท่อที่ 1: สำหรับคุณแม่ (แยกบัญชีขาด)", "เงินปันผลหุ้น สอ.กฟผ. (เงินต้น 4.0M @ 4.5%)", 180000, 15000, "โอนตรงให้คุณแม่ดูแลค่าใช้จ่ายถึงอายุ 98 ปี"),
    ("ท่อที่ 2: ดอกเบี้ยพอร์ตส่วนตัว", "ผลตอบแทนจาก 5 ตะกร้าที่เหลือ (เงินต้น 21.0M)", 468050, 39004, "ผลิตดอกเบี้ยสม่ำเสมอ เงินต้น 21M อยู่ครบ 100%"),
    ("ท่อที่ 2: บำนาญการันตี 1", "บำนาญเดิม (เมืองไทยประกันชีวิต)", 70000, 5833, "การันตีรับถึงอายุ 85 ปี"),
    ("ท่อที่ 2: บำนาญการันตี 2", "บำนาญใหม่ #3 (จ่ายสั้นช่วงอายุ 50-54)", 32400, 2700, "การันตีรับถึงอายุ 85 ปี"),
    ("ท่อที่ 2: บำนาญการันตี 3", "บำนาญ 85/5 (จ่ายสั้นช่วงอายุ 53-57)", 30000, 2500, "การันตีรับถึงอายุ 85 ปี")
]

for idx, item in enumerate(cf_pipeline_data, start=7):
    pipe, src, yr, mo, obj = item
    ws3.row_dimensions[idx].height = 22
    ws3.cell(row=idx, column=2, value=pipe).font = bold_font
    ws3.cell(row=idx, column=3, value=src).font = normal_font
    ws3.cell(row=idx, column=4, value=yr).number_format = "#,##0"
    ws3.cell(row=idx, column=4).alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(row=idx, column=5, value=mo).number_format = "#,##0"
    ws3.cell(row=idx, column=5).alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(row=idx, column=5).font = bold_font
    ws3.cell(row=idx, column=6, value=obj).font = normal_font
    
    for c in range(2, 7):
        ws3.cell(row=idx, column=c).border = thin_border
        ws3.cell(row=idx, column=c).fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

tot_pipe_row = 12
ws3.row_dimensions[tot_pipe_row].height = 26
ws3.cell(row=tot_pipe_row, column=2, value="รวมกระแสเงินสดรับส่วนตัวสุทธิ (ท่อที่ 2)").font = bold_font
ws3.cell(row=tot_pipe_row, column=3, value="ดอกเบี้ยพอร์ต 21M + บำนาญ 3 ฉบับ (11,033 บ./ด.)").font = normal_font
ws3.cell(row=tot_pipe_row, column=4, value=f"=SUM(D8:D11)").number_format = "#,##0"
ws3.cell(row=tot_pipe_row, column=5, value=f"=SUM(E8:E11)").number_format = "#,##0"
ws3.cell(row=tot_pipe_row, column=6, value="เงินสดรับส่วนตัว 50,038 บาท/เดือน (600,450 บ./ปี)").font = bold_font

for c in range(2, 7):
    cell = ws3.cell(row=tot_pipe_row, column=c)
    cell.font = bold_font
    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    cell.border = total_border
    if c in [4, 5]:
        cell.alignment = Alignment(horizontal="right", vertical="center")

# Section 2: Personal Budget Breakdown
ws3["B14"] = "2. การจัดสรรงบประมาณส่วนตัวรายเดือน (Personal Expense & Medical Buffer)"
ws3["B14"].font = Font(name="Arial", size=12, bold=True, color=NAVY_HEADER)

headers3_2 = [
    ("รายการรายจ่าย", 26),
    ("สัดส่วนในกระแสเงินสด", 34),
    ("งบประมาณต่อปี (บาท)", 20),
    ("งบประมาณต่อเดือน (บาท)", 22),
    ("คำอธิบายการบริหารจัดการ", 40)
]
for col_idx, (h_text, width) in enumerate(headers3_2, start=2):
    cell = ws3.cell(row=15, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = PatternFill(start_color=EMERALD_HEADER, end_color=EMERALD_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[15].height = 28

budget_data = [
    ("งบกินอยู่ใช้จ่ายในชีวิตประจำวัน", "69.9% ของกระแสเงินสดรับส่วนตัว", 420000, 35000, "ใช้จ่ายอิสระตามไลฟ์สไตล์ เกษียณอย่างมีคุณภาพ"),
    ("งบส่วนเกินสำหรับเบี้ยประกันสุขภาพ (Surplus)", "30.1% ของกระแสเงินสดรับส่วนตัว", 180450, 15038, "รองรับเบี้ยสุขภาพที่เพิ่มขึ้นตามอายุ โดยไม่แตะเงินต้น 25M")
]

for idx, item in enumerate(budget_data, start=16):
    ws3.row_dimensions[idx].height = 22
    name, ratio, yr, mo, desc = item
    ws3.cell(row=idx, column=2, value=name).font = bold_font
    ws3.cell(row=idx, column=3, value=ratio).font = normal_font
    ws3.cell(row=idx, column=4, value=yr).number_format = "#,##0"
    ws3.cell(row=idx, column=4).alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(row=idx, column=5, value=mo).number_format = "#,##0"
    ws3.cell(row=idx, column=5).alignment = Alignment(horizontal="right", vertical="center")
    ws3.cell(row=idx, column=5).font = bold_font
    ws3.cell(row=idx, column=6, value=desc).font = normal_font
    for c in range(2, 7):
        ws3.cell(row=idx, column=c).border = thin_border
        ws3.cell(row=idx, column=c).fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

# Section 3: Drawdown Protocol
ws3["B20"] = "3. ลำดับขั้นตอนการทยอยดึงเงินต้นมาใช้ (Drawdown Execution Protocol)"
ws3["B20"].font = Font(name="Arial", size=12, bold=True, color=NAVY_HEADER)

drawdown_headers = [
    ("ช่วงอายุ (ปี)", 16),
    ("ระยะกลยุทธ์", 24),
    ("สถานะเงินต้น 25 ล้านบาท", 26),
    ("การบริหารจัดการสินทรัพย์", 45)
]
for col_idx, (h_text, width) in enumerate(drawdown_headers, start=2):
    cell = ws3.cell(row=21, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[21].height = 26

drawdown_data = [
    ("60 - 70 ปี", "Capital Preservation Phase", "เงินต้นคงอยู่ครบ 100% (25.00M)", "ใช้จ่ายเฉพาะดอกเบี้ยและบำนาญ 50,038 บ./ด. แบ่ง 500k จากตะกร้า 6 ไว้เป็นกองทุนฉุกเฉินคุณแม่"),
    ("71 - 79 ปี", "Tier-1 Drawdown Phase", "เริ่มทยอยดึงสลากออมทรัพย์ 3.0M", "เบี้ยสุขภาพแตะ 180k-250k บ./ปี ให้ไถ่ถอนสลากออมทรัพย์ปีละ 300k-350k + เงินคืนสะสมทรัพย์ 250k (ตอนอายุ 62)"),
    ("80 ปีขึ้นไป", "Core Liquidation Phase", "ทยอยปลดล็อก MMF และพันธบัตร", "คุณแม่อายุ 98 ปี หุ้น สอ. 4M ปลดล็อกกลับมาเป็นสภาพคล่องส่วนตัว ไถ่ถอน MMF 4M และพันธบัตร 6.5M ตามลำดับ โดยคงเงินฝากพิเศษ 5M และทองคำ 5 บาทไว้ท้ายสุด")
]

for idx, item in enumerate(drawdown_data, start=22):
    ws3.row_dimensions[idx].height = 24
    age_range, phase, cap_status, details = item
    ws3.cell(row=idx, column=2, value=age_range).alignment = Alignment(horizontal="center", vertical="center")
    ws3.cell(row=idx, column=2).font = bold_font
    ws3.cell(row=idx, column=3, value=phase).font = bold_font
    ws3.cell(row=idx, column=4, value=cap_status).font = normal_font
    ws3.cell(row=idx, column=5, value=details).font = normal_font
    for c in range(2, 6):
        ws3.cell(row=idx, column=c).border = thin_border
        ws3.cell(row=idx, column=c).fill = PatternFill(start_color=ZEBRA_BG if idx%2==0 else WHITE, end_color=ZEBRA_BG if idx%2==0 else WHITE, fill_type="solid")

output_file = "/Users/tor/retire-plan/Retirement_Plan_25M.xlsx"
wb.save(output_file)
print(f"Updated Excel saved at {output_file}")

import csv

csv_file = "/Users/tor/retire-plan/Retirement_Plan.csv"
with open(csv_file, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerow([
        "อายุ", "ฐานเงินเดือน", "โบนัสประจำปี(~2ด.)", "รวมรายได้ต่อปี",
        "PVD_15%", "สหกรณ์_สอกฟผ", "RMF", "ประกันเดิม_3ฉบับ", "ฝากพิเศษ_สอกฟผ",
        "รวมหักจากเงินเดือนประจำ", "บำนาญใหม่_หักจากโบนัสรายปี", "รวมจ่ายทั้งหมดต่อปี",
        "เงินเหลือใช้สุทธิ/ปี", "หมุดหมายสำคัญ"
    ])
    for r in raw_data_s1:
        seq, age, salary, coop, rmf, ins_orig, coop_spec, bonus_pen, note = r
        bonus = salary * 2
        total_income = salary * 12 + bonus
        pvd = round(salary * 0.15)
        monthly_deduct = pvd + coop + rmf + ins_orig + coop_spec
        total_paid = monthly_deduct * 12 + bonus_pen
        net_disposable = total_income - total_paid
        writer.writerow([
            age, salary, bonus, total_income,
            pvd, coop, rmf, ins_orig, coop_spec,
            monthly_deduct, bonus_pen, total_paid,
            net_disposable, note
        ])
print(f"Updated CSV saved at {csv_file}")

